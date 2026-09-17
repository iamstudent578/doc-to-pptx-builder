"""
add_speaker_notes.py
=====================
把步驟 5 生成的逐頁講稿批次寫入 pptx 的備忘稿欄位。

用法：
    python add_speaker_notes.py deck.pptx notes.json output.pptx

notes.json 格式：
{
  "notes": [
    "第 1 頁的講稿全文...",
    "第 2 頁的講稿全文...",
    ...
  ]
}
陣列順序對應投影片順序（第 0 項 = 第 1 頁），長度必須等於投影片總數，
如果某頁沒有講稿就放空字串 ""，不要省略該項目，避免後面全部對不齊。
"""

import json
import sys

from pptx import Presentation


def add_notes(pptx_path: str, notes_json_path: str, output_path: str):
    with open(notes_json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    notes = data["notes"]

    prs = Presentation(pptx_path)

    if len(notes) != len(prs.slides):
        raise ValueError(
            f"講稿數量 ({len(notes)}) 跟投影片數量 ({len(prs.slides)}) 對不上，"
            "請檢查 notes.json 是否每頁都有對應項目（沒講稿的頁面請放空字串而不是省略）。"
        )

    for slide, note_text in zip(prs.slides, notes):
        if not note_text:
            continue
        notes_slide = slide.notes_slide  # 若不存在會自動建立
        notes_slide.notes_text_frame.text = note_text

    prs.save(output_path)
    print(f"已寫入 {len(notes)} 頁備忘稿，輸出檔案: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print("用法: python add_speaker_notes.py deck.pptx notes.json output.pptx")
        sys.exit(1)
    add_notes(sys.argv[1], sys.argv[2], sys.argv[3])
