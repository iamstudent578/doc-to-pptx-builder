"""
extract_docs.py
================
把 .docx / .pdf 文件解析成結構化的「內容清單」JSON，供步驟 2 擬定簡報大綱使用。

v3 更新重點（對應精簡版 SKILL.md 的 Embedded Assets 要求）：
- 圖片依照它在文件中實際出現的位置擷取，不再整批掛在最後一個 section。
- 每張圖片附帶 context（前後文）與 caption（文件原生圖說/替代文字，沒有則為 null）。
- 新增 used 欄位（預設 false），供 Step 3 Asset Matching 配對成功後標記，避免同一張圖
  被配到兩張投影片。
- 刻意不做 asset_id / order_index / asset_type / auto_caption 這類欄位——這個精簡版本
  的配對完全交給 agent 自己讀 context 判斷，不需要靠額外欄位做評分或分類。

設計原則：寧可讓後續步驟看到「稍微多一點原始資訊」，也不要在這一步就過度摘要，
因為擬大綱時才是真正該做取捨判斷的地方。

依賴（依實際環境安裝，找不到就 pip install）：
    pip install python-docx pymupdf

輸出格式範例：
{
  "source_file": "report.docx",
  "sections": [
    {
      "heading": "市場現況",
      "level": 1,
      "paragraphs": ["...段落原文，供後續濃縮重點用..."],
      "key_numbers": [{"value": "35%", "context": "去年成長率為 35%"}],
      "tables": [{"caption": null, "rows": [["A", "B"], ["1", "2"]]}],
      "images": [
        {
          "path": "extracted_images/img_0001.png",
          "context": "前文：去年市場成長顯著。 後文：預期今年將延續此趨勢。",
          "caption": "圖 1：市場成長趨勢",
          "used": false
        }
      ]
    }
  ]
}
"""

import json
import os
import re
import sys

NUMBER_PATTERN = re.compile(r"\d[\d,.]*%?")
CAPTION_PATTERN = re.compile(r"^(圖|表|Figure|Fig\.|Table)\s*\d*", re.IGNORECASE)


def _new_section(heading: str, level: int = 1) -> dict:
    return {
        "heading": heading,
        "level": level,
        "paragraphs": [],
        "key_numbers": [],
        "tables": [],
        "images": [],
    }


def _build_context(before: str | None, after: str | None) -> str:
    parts = []
    if before:
        parts.append(f"前文：{before}")
    if after:
        parts.append(f"後文：{after}")
    return " ".join(parts)


def extract_docx(path: str, image_out_dir: str) -> dict:
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(path)
    os.makedirs(image_out_dir, exist_ok=True)
    rels = doc.part.rels

    body_paragraphs = doc.paragraphs  # 已依文件順序排列
    sections = []
    current = None
    img_counter = 0

    def para_text(p_obj) -> str:
        return p_obj.text.strip()

    def para_style(p_obj) -> str:
        return (p_obj.style.name or "").lower()

    for idx, p in enumerate(body_paragraphs):
        style = para_style(p)
        text = para_text(p)

        # 標題 -> 開新 section
        if style.startswith("heading") and text:
            level_match = re.search(r"\d+", style)
            level = int(level_match.group()) if level_match else 1
            current = _new_section(text, level)
            sections.append(current)
            continue

        if current is None:
            current = _new_section("(無標題章節)", 1)
            sections.append(current)

        # 偵測此段落內是否包含內嵌圖片（drawing/blip）
        blips = p._p.findall(".//" + qn("a:blip"))
        docpr_list = p._p.findall(".//" + qn("wp:docPr"))

        if blips:
            for i, blip in enumerate(blips):
                r_id = blip.get(qn("r:embed"))
                if not r_id or r_id not in rels:
                    continue
                rel = rels[r_id]
                img_counter += 1
                ext = os.path.splitext(rel.target_ref)[1] or ".png"
                out_path = os.path.join(image_out_dir, f"img_{img_counter:04d}{ext}")
                with open(out_path, "wb") as f:
                    f.write(rel.target_part.blob)

                # alt text（若文件作者有填替代文字，視為 caption 候選）
                alt_text = None
                if i < len(docpr_list):
                    alt_text = docpr_list[i].get("descr") or docpr_list[i].get("title") or None
                    if alt_text:
                        alt_text = alt_text.strip() or None

                # 前文：目前 section 已收集到的最後一段文字，沒有就退回標題
                if current["paragraphs"]:
                    anchor_before = current["paragraphs"][-1]
                elif current["heading"] != "(無標題章節)":
                    anchor_before = current["heading"]
                else:
                    anchor_before = None

                # 後文：往下找最近 1-2 個非空段落，若像圖說格式則視為 caption
                anchor_after = None
                caption = alt_text
                for j in range(idx + 1, min(idx + 3, len(body_paragraphs))):
                    next_text = para_text(body_paragraphs[j])
                    next_style = para_style(body_paragraphs[j])
                    if not next_text:
                        continue
                    if caption is None and (
                        "caption" in next_style or CAPTION_PATTERN.match(next_text)
                    ):
                        caption = next_text
                        continue
                    anchor_after = next_text
                    break

                current["images"].append({
                    "path": out_path,
                    "context": _build_context(anchor_before, anchor_after),
                    "caption": caption,
                    "used": False,
                })
            # 圖片段落本身通常沒有其他有意義文字，跳過文字收集
            continue

        if text:
            current["paragraphs"].append(text)
            for match in NUMBER_PATTERN.finditer(text):
                current["key_numbers"].append({"value": match.group(), "context": text})

    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        target = sections[-1] if sections else None
        if target:
            target["tables"].append({"caption": None, "rows": rows})

    return {"source_file": path, "sections": sections}


def extract_pdf(path: str, image_out_dir: str) -> dict:
    import fitz  # PyMuPDF

    os.makedirs(image_out_dir, exist_ok=True)
    pdf = fitz.open(path)
    sections = []
    img_counter = 0

    for page_index in range(len(pdf)):
        page = pdf[page_index]
        raw_text = page.get_text().strip()
        images_on_page = page.get_images(full=True)
        if not raw_text and not images_on_page:
            continue

        # blocks: (x0, y0, x1, y1, text, block_no, block_type)，由上到下排序供 anchor 判斷
        text_blocks = [b for b in page.get_text("blocks") if b[4].strip()]
        text_blocks.sort(key=lambda b: b[1])

        section = _new_section(f"第 {page_index + 1} 頁", 1)
        section["paragraphs"] = [b[4].strip() for b in text_blocks]
        for line in section["paragraphs"]:
            for match in NUMBER_PATTERN.finditer(line):
                section["key_numbers"].append({"value": match.group(), "context": line})

        for img in images_on_page:
            xref = img[0]
            img_counter += 1
            base = pdf.extract_image(xref)
            out_path = os.path.join(image_out_dir, f"img_{img_counter:04d}.{base['ext']}")
            with open(out_path, "wb") as f:
                f.write(base["image"])

            rects = page.get_image_rects(xref)
            img_rect = rects[0] if rects else None

            anchor_before, anchor_after, caption = None, None, None
            if img_rect is not None:
                above = [b for b in text_blocks if b[3] <= img_rect.y0]
                below = [b for b in text_blocks if b[1] >= img_rect.y1]
                if above:
                    anchor_before = above[-1][4].strip()
                if below:
                    first_below = below[0][4].strip()
                    if CAPTION_PATTERN.match(first_below):
                        caption = first_below
                        if len(below) > 1:
                            anchor_after = below[1][4].strip()
                    else:
                        anchor_after = first_below

            section["images"].append({
                "path": out_path,
                "context": _build_context(anchor_before, anchor_after),
                "caption": caption,
                "used": False,
            })

        sections.append(section)

    # 掃描版 PDF（每頁幾乎沒有文字但有圖）需要另外接 OCR，這裡先不處理，
    # 若 sections 內容多為空但頁數很多，代表可能是掃描件，應提醒使用者確認。
    return {"source_file": path, "sections": sections}


def main():
    if len(sys.argv) < 2:
        print("用法: python extract_docs.py <file.docx|file.pdf> [輸出json路徑]")
        sys.exit(1)

    src = sys.argv[1]
    out_json = sys.argv[2] if len(sys.argv) > 2 else "content_inventory.json"
    image_dir = os.path.join(os.path.dirname(out_json) or ".", "extracted_images")

    ext = os.path.splitext(src)[1].lower()
    if ext == ".docx":
        result = extract_docx(src, image_dir)
    elif ext == ".pdf":
        result = extract_pdf(src, image_dir)
    else:
        raise ValueError(f"不支援的檔案類型: {ext}")

    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    total_images = sum(len(s["images"]) for s in result["sections"])
    print(f"已輸出內容清單: {out_json}（共 {len(result['sections'])} 個 section, {total_images} 張圖片資產）")


if __name__ == "__main__":
    main()
