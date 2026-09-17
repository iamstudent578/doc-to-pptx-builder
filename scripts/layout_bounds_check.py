#!/usr/bin/env python3
"""
layout_bounds_check.py

Deterministic, structural QA check for generated .pptx files, used as the
Step 4 fallback when LibreOffice headless rendering is unavailable (so no
pixel-level contact-sheet review is possible).

This script intentionally does NOT try to be clever or heuristic-heavy.
It checks only things that can be verified from shape geometry and text
properties, and it explicitly excludes known-safe patterns (like a
background rectangle fully containing a text box) so it doesn't generate
false-positive "overlap" noise.

Usage:
    python3 layout_bounds_check.py <path_to_pptx> [--json-out report.json]

Exit code:
    0 -> no issues found
    1 -> issues found (see stdout / report.json)

Checks performed:
    1. Shapes positioned off-slide (negative or beyond slide bounds).
    2. Font sizes outside a sane readable range (default 10-60pt).
    3. Paragraphs whose estimated wrapped line count exceeds the
       text box's available height (rough overflow estimate).
    4. Shape-pair overlaps, EXCLUDING:
         - decorative hairlines (very thin shapes)
         - pure containment (one shape's box is ~fully inside another's,
           e.g. a text box sitting on top of a background card/rectangle)

What this script deliberately does NOT do:
    - It does not render pixels or judge color contrast, alignment
      aesthetics, or image quality. That requires actual visual review
      (LibreOffice + vision model). If that path is available, prefer it.
"""

import argparse
import json
import sys
from dataclasses import dataclass, field

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:
    print("ERROR: python-pptx is required (pip install python-pptx --break-system-packages)")
    sys.exit(2)

EMU_PER_INCH = 914400

# Tunable thresholds
MIN_FONT_PT = 10
MAX_FONT_PT = 60
MAX_CHARS_PER_LINE_ESTIMATE = 40   # rough CJK/Latin-mixed default
LINE_HEIGHT_INCHES_ESTIMATE = 0.35
MAX_PARAGRAPH_CHARS = 200
CONTAINMENT_THRESHOLD = 0.95       # >=95% of smaller box area inside larger = containment, not overlap
THIN_SHAPE_THRESHOLD_IN = 0.05     # shapes thinner than this on either axis are treated as decorative lines
BOUNDARY_TOLERANCE_IN = 0.05       # small tolerance before flagging off-slide


def emu_to_in(v):
    return v / EMU_PER_INCH


@dataclass
class ShapeInfo:
    name: str
    left: float
    top: float
    width: float
    height: float
    right: float = field(init=False)
    bottom: float = field(init=False)
    has_text: bool = False

    def __post_init__(self):
        self.right = self.left + self.width
        self.bottom = self.top + self.height

    def area(self):
        return max(self.width, 0) * max(self.height, 0)

    def intersection_area(self, other: "ShapeInfo"):
        ix = max(0.0, min(self.right, other.right) - max(self.left, other.left))
        iy = max(0.0, min(self.bottom, other.bottom) - max(self.top, other.top))
        return ix * iy

    def is_thin(self):
        return self.width < THIN_SHAPE_THRESHOLD_IN or self.height < THIN_SHAPE_THRESHOLD_IN


def check_slide(slide, slide_num, slide_w_in, slide_h_in):
    issues = []
    shapes_info = []

    for shape in slide.shapes:
        try:
            left = emu_to_in(shape.left) if shape.left is not None else 0.0
            top = emu_to_in(shape.top) if shape.top is not None else 0.0
            width = emu_to_in(shape.width) if shape.width is not None else 0.0
            height = emu_to_in(shape.height) if shape.height is not None else 0.0
        except (TypeError, AttributeError):
            continue

        info = ShapeInfo(name=shape.name, left=left, top=top, width=width, height=height)
        info.has_text = bool(getattr(shape, "has_text_frame", False) and shape.has_text_frame)
        shapes_info.append(info)

        # 1. Off-slide check
        if left < -BOUNDARY_TOLERANCE_IN or top < -BOUNDARY_TOLERANCE_IN:
            issues.append({
                "type": "off_slide",
                "shape": shape.name,
                "detail": f"negative position left={left:.2f}in top={top:.2f}in",
            })
        if info.right > slide_w_in + BOUNDARY_TOLERANCE_IN or info.bottom > slide_h_in + BOUNDARY_TOLERANCE_IN:
            issues.append({
                "type": "off_slide",
                "shape": shape.name,
                "detail": f"exceeds slide bounds right={info.right:.2f}in bottom={info.bottom:.2f}in "
                          f"(slide is {slide_w_in:.2f}x{slide_h_in:.2f}in)",
            })

        # 2 & 3. Font size + overflow estimate
        if info.has_text:
            tf = shape.text_frame
            total_chars = 0
            for p in tf.paragraphs:
                p_text = "".join(r.text for r in p.runs) or p.text
                total_chars += len(p_text)

                if len(p_text) > MAX_PARAGRAPH_CHARS:
                    issues.append({
                        "type": "long_paragraph",
                        "shape": shape.name,
                        "detail": f"{len(p_text)} chars: '{p_text[:50]}...'",
                    })

                # font size: check paragraph-level and run-level
                sizes_pt = []
                if p.font.size is not None:
                    sizes_pt.append(p.font.size.pt)
                for r in p.runs:
                    if r.font.size is not None:
                        sizes_pt.append(r.font.size.pt)

                for fs in sizes_pt:
                    if fs < MIN_FONT_PT:
                        issues.append({
                            "type": "font_too_small",
                            "shape": shape.name,
                            "detail": f"{fs:.0f}pt on '{p_text[:30]}'",
                        })
                    if fs > MAX_FONT_PT:
                        issues.append({
                            "type": "font_too_large",
                            "shape": shape.name,
                            "detail": f"{fs:.0f}pt on '{p_text[:30]}'",
                        })

            if total_chars > 0 and height > 0:
                est_lines = max(1, total_chars / MAX_CHARS_PER_LINE_ESTIMATE)
                est_needed_height = est_lines * LINE_HEIGHT_INCHES_ESTIMATE
                if est_needed_height > height * 1.05:  # small tolerance
                    issues.append({
                        "type": "possible_text_overflow",
                        "shape": shape.name,
                        "detail": f"~{est_lines:.0f} estimated lines need ~{est_needed_height:.2f}in, "
                                  f"box height is {height:.2f}in",
                    })

    # 4. Overlap check, excluding containment and decorative thin shapes
    for i, a in enumerate(shapes_info):
        for b in shapes_info[i + 1:]:
            if a.is_thin() or b.is_thin():
                continue
            if a.area() <= 0 or b.area() <= 0:
                continue

            inter = a.intersection_area(b)
            if inter <= 0:
                continue

            smaller_area = min(a.area(), b.area())
            containment_ratio = inter / smaller_area if smaller_area > 0 else 0

            if containment_ratio >= CONTAINMENT_THRESHOLD:
                # One shape sits (almost) fully inside/on top of the other —
                # e.g. text on a background card. This is expected, not an issue.
                continue

            issues.append({
                "type": "unexpected_overlap",
                "shapes": [a.name, b.name],
                "detail": f"overlap area {inter:.2f}in^2, containment_ratio={containment_ratio:.2f}",
            })

    return {
        "slide": slide_num,
        "issue_count": len(issues),
        "issues": issues,
    }


def run_check(pptx_path):
    prs = Presentation(pptx_path)
    slide_w_in = emu_to_in(prs.slide_width)
    slide_h_in = emu_to_in(prs.slide_height)

    report = {
        "file": pptx_path,
        "slide_width_in": round(slide_w_in, 3),
        "slide_height_in": round(slide_h_in, 3),
        "slide_count": len(prs.slides._sldIdLst),
        "slides": [],
        "total_issue_count": 0,
    }

    for idx, slide in enumerate(prs.slides, start=1):
        slide_report = check_slide(slide, idx, slide_w_in, slide_h_in)
        report["slides"].append(slide_report)
        report["total_issue_count"] += slide_report["issue_count"]

    return report


def print_human_readable(report):
    print(f"Checked: {report['file']}")
    print(f"Slide size: {report['slide_width_in']}x{report['slide_height_in']} in, "
          f"{report['slide_count']} slides\n")

    if report["total_issue_count"] == 0:
        print("✅ No structural issues found.")
        return

    for slide_report in report["slides"]:
        if slide_report["issue_count"] == 0:
            continue
        print(f"Slide {slide_report['slide']}: {slide_report['issue_count']} issue(s)")
        for issue in slide_report["issues"]:
            shape_ref = issue.get("shape") or issue.get("shapes")
            print(f"  - [{issue['type']}] {shape_ref}: {issue['detail']}")
        print()

    print(f"⚠ Total issues: {report['total_issue_count']}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx_path", help="Path to the .pptx file to check")
    parser.add_argument("--json-out", help="Optional path to write the full JSON report")
    args = parser.parse_args()

    report = run_check(args.pptx_path)
    print_human_readable(report)

    if args.json_out:
        with open(args.json_out, "w", encoding="utf-8") as f:
            json.dump(report, f, ensure_ascii=False, indent=2)
        print(f"\nFull JSON report written to {args.json_out}")

    sys.exit(1 if report["total_issue_count"] > 0 else 0)


if __name__ == "__main__":
    main()
